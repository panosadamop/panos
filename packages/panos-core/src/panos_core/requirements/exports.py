"""Office/PDF exports of a project's requirements (FR-029, BR-028, M8-05).

* ``SPECIFICATION.docx`` / ``.pdf``: every Markdown document (BRD … risks) in one file, each
  starting on a new page, rendered from the same Markdown as ``docs/requirements``.
* ``REQUIREMENTS.xlsx``: requirement register, backlog with estimates, cost by role, sprint
  plan and trace links, as filterable sheets.
* ``SUMMARY.pptx``: a short overview deck (scope, stack, estimate, risks, open decisions).

Like the Markdown documents, exports are derived output: regenerate, don't edit.
"""

import io
from collections import Counter
from collections.abc import Iterable
from typing import Any, Final

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter
from pptx import Presentation
from pptx.util import Inches, Pt

from panos_core.documents import FontSet, render_docx, render_pdf
from panos_core.domain.requirements import RequirementStatus, RequirementType
from panos_core.requirements.artifacts import DOCUMENTS, ArtifactContext, render_all

FORMATS: Final = ("docx", "pdf", "xlsx", "pptx")
_HEADER_FILL: Final = PatternFill("solid", fgColor="DDE6F0")


class ExportError(ValueError):
    pass


def _specification_markdown(c: ArtifactContext) -> str:
    rendered = render_all(c)
    return "\n\n".join(rendered[name] for name in DOCUMENTS)


def _sheet(wb: Workbook, title: str, header: list[str], rows: Iterable[list[Any]]) -> None:
    ws = wb.create_sheet(title)
    ws.append(header)
    for cell in ws[1]:
        cell.font = Font(bold=True)
        cell.fill = _HEADER_FILL
    widths = [len(h) for h in header]
    for row in rows:
        ws.append(row)
        widths = [
            max(w, min(60, len(str(v)) if v is not None else 0))
            for w, v in zip(widths, row, strict=False)
        ]
    for i, width in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(i)].width = width + 2
    ws.freeze_panes = "A2"
    if ws.max_row > 1:
        ws.auto_filter.ref = ws.dimensions


def workbook(c: ArtifactContext) -> bytes:
    wb = Workbook()
    for default in list(wb.worksheets):  # start without the default empty sheet
        wb.remove(default)
    items = sorted(c.items, key=lambda i: i.key)
    _sheet(
        wb,
        "Requirements",
        [
            "Key",
            "Type",
            "Title",
            "Version",
            "Status",
            "Stale",
            "Parent",
            "Priority",
            "Points",
            "Source",
        ],
        (
            [
                i.key,
                i.type.value,
                i.title,
                i.version,
                i.status.value,
                "yes" if i.stale else "",
                i.parent_key or "",
                str(i.attributes.get("priority", "")),
                i.attributes.get("points", ""),
                ", ".join(i.source),
            ]
            for i in items
        ),
    )
    est = c.estimate
    if est is not None:
        _sheet(
            wb,
            "Backlog",
            ["Story", "Title", "Points", "Estimated", "Priority", "Hours", "Person-days"],
            (
                [
                    s.key,
                    s.title,
                    s.points,
                    "yes" if s.estimated else "default",
                    s.priority,
                    s.hours,
                    s.person_days,
                ]
                for s in est.stories
            ),
        )
        _sheet(
            wb,
            "Cost",
            ["Role", "Hours", f"Rate ({est.currency})", f"Cost ({est.currency})"],
            [
                *([role, v["hours"], v["rate"], v["cost"]] for role, v in est.by_role.items()),
                ["Total", est.total_hours, None, est.total_cost],
            ],
        )
        _sheet(
            wb,
            "Sprints",
            ["Sprint", "Points", "Stories"],
            ([s.number, s.points, ", ".join(s.stories)] for s in est.sprints),
        )
    _sheet(
        wb,
        "Trace links",
        ["Source", "Relation", "Target"],
        (
            [link.source_key, link.relation.value, link.target]
            for link in sorted(c.links, key=lambda x: (x.source_key, x.target))
        ),
    )
    wb.properties.title = f"{c.project} — requirements"
    wb.properties.creator = "Panos"
    buffer = io.BytesIO()
    wb.save(buffer)
    return buffer.getvalue()


def _bullets(prs: Any, title: str, lines: list[str], *, empty: str = "None.") -> None:
    slide = prs.slides.add_slide(prs.slide_layouts[1])
    slide.shapes.title.text = title
    body = slide.placeholders[1].text_frame
    lines = lines or [empty]
    body.text = lines[0]
    for line in lines[1:12]:
        body.add_paragraph().text = line
    if len(lines) > 12:
        body.add_paragraph().text = f"… and {len(lines) - 12} more"
    for paragraph in body.paragraphs:
        for run in paragraph.runs:
            run.font.size = Pt(18 if len(lines) <= 6 else 14)


def deck(c: ArtifactContext) -> bytes:
    prs = Presentation()
    prs.slide_width, prs.slide_height = Inches(13.333), Inches(7.5)
    items = sorted(c.items, key=lambda i: i.key)
    by_type = Counter(i.type.value for i in items)
    approved = sum(1 for i in items if i.status == RequirementStatus.APPROVED)

    title = prs.slides.add_slide(prs.slide_layouts[0])
    title.shapes.title.text = c.project
    title.placeholders[
        1
    ].text = f"Requirements summary · {approved}/{len(items)} items approved · generated by Panos"

    of = lambda t: [i for i in items if i.type == t]  # noqa: E731
    _bullets(prs, "Business goals", [f"{i.key} — {i.title}" for i in of(RequirementType.BR)])
    scope = []
    for epic in of(RequirementType.EPIC):
        scope.append(f"{epic.key} — {epic.title}")
        scope += [
            f"    {f.key} — {f.title}" for f in of(RequirementType.FEAT) if f.parent_key == epic.key
        ]
    _bullets(
        prs,
        "Scope",
        scope or [f"{k}: {v}" for k, v in sorted(by_type.items())],
    )
    if c.stack:
        stack = [f"{c.stack.get('name', c.stack.get('pack'))}"]
        stack += [f"Add-on: {a}" for a in c.stack.get("addons", [])]
        _bullets(prs, "Technology stack", stack)
    est = c.estimate
    if est is not None:
        _bullets(
            prs,
            "Estimate",
            [
                f"{est.total_points} story points · {est.total_hours:.1f} h"
                f" · {est.person_days:.1f} person-days",
                f"Cost: {est.total_cost:,.2f} {est.currency}",
                f"Confidence: {est.confidence} ({est.confidence_score:.2f})",
                f"Sprints: {len(est.sprints)}",
                *est.assumptions,
            ],
        )
    _bullets(prs, "Risks", [f"{i.key} — {i.title}" for i in of(RequirementType.RISK)])
    pending = [
        f"{i.key} — {i.title} ({'STALE' if i.stale else i.status.value})"
        for i in items
        if i.stale or i.status in (RequirementStatus.DRAFT, RequirementStatus.PENDING_APPROVAL)
    ]
    _bullets(prs, "Open decisions", pending, empty="Everything is approved and current.")
    prs.core_properties.title = f"{c.project} — summary"
    prs.core_properties.author = "Panos"
    buffer = io.BytesIO()
    prs.save(buffer)
    return buffer.getvalue()


def export(
    c: ArtifactContext, formats: Iterable[str] = FORMATS, *, fonts: FontSet | None = None
) -> dict[str, bytes]:
    """File name → bytes for the requested formats."""
    wanted = list(dict.fromkeys(f.lower().lstrip(".") for f in formats))
    unknown = [f for f in wanted if f not in FORMATS]
    if unknown or not wanted:
        raise ExportError(f"unknown export format(s) {unknown}; choose from {', '.join(FORMATS)}")
    out: dict[str, bytes] = {}
    title = f"{c.project} — Specification"
    if "docx" in wanted or "pdf" in wanted:
        markdown = _specification_markdown(c)
        if "docx" in wanted:
            out["SPECIFICATION.docx"] = render_docx(markdown, title=title)
        if "pdf" in wanted:
            out["SPECIFICATION.pdf"] = render_pdf(markdown, title=title, fonts=fonts)
    if "xlsx" in wanted:
        out["REQUIREMENTS.xlsx"] = workbook(c)
    if "pptx" in wanted:
        out["SUMMARY.pptx"] = deck(c)
    return out

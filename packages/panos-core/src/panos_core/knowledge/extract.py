"""Text extraction from Office and PDF files for the knowledge index (FR-019, M8-04).

The extracted text keeps location markers (``[page 3]``, ``[slide 2]``, ``[sheet Orders]``)
so a hit's line range still points somewhere a person can find. Extraction never executes
anything: macros, embedded objects and scripts are ignored. A damaged or encrypted file
yields ``None`` and is skipped, never an error for the whole index run.
"""

import io
from collections.abc import Callable
from pathlib import PurePosixPath
from typing import Final

from panos_core.observability import get_logger

log = get_logger(__name__)

MAX_DOCUMENT_BYTES: Final = 25_000_000  # Office/PDF files (text files keep the 1 MB limit)
MAX_EXTRACTED_CHARS: Final = 2_000_000
MAX_SHEET_ROWS: Final = 5_000


def _docx(data: bytes) -> str:
    from docx import Document

    doc = Document(io.BytesIO(data))
    lines = [p.text for p in doc.paragraphs]
    for n, table in enumerate(doc.tables, start=1):
        lines.append(f"[table {n}]")
        lines += [" | ".join(cell.text.strip() for cell in row.cells) for row in table.rows]
    return "\n".join(lines)


def _xlsx(data: bytes) -> str:
    from openpyxl import load_workbook

    wb = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    lines: list[str] = []
    try:
        for ws in wb.worksheets:
            lines.append(f"[sheet {ws.title}]")
            for i, row in enumerate(ws.iter_rows(values_only=True)):
                if i >= MAX_SHEET_ROWS:
                    lines.append(f"[sheet {ws.title} truncated after {MAX_SHEET_ROWS} rows]")
                    break
                cells = ["" if v is None else str(v) for v in row]
                if any(cells):
                    lines.append("\t".join(cells).rstrip())
    finally:
        wb.close()
    return "\n".join(lines)


def _pptx(data: bytes) -> str:
    from pptx import Presentation

    prs = Presentation(io.BytesIO(data))
    lines: list[str] = []
    for n, slide in enumerate(prs.slides, start=1):
        lines.append(f"[slide {n}]")
        for shape in slide.shapes:
            if shape.has_text_frame:
                lines += [p.text for p in shape.text_frame.paragraphs if p.text.strip()]
            if getattr(shape, "has_table", False):
                for row in shape.table.rows:
                    lines.append(" | ".join(c.text.strip() for c in row.cells))
        if slide.has_notes_slide and slide.notes_slide.notes_text_frame is not None:
            notes = slide.notes_slide.notes_text_frame.text.strip()
            if notes:
                lines.append(f"[notes] {notes}")
    return "\n".join(lines)


def _pdf(data: bytes) -> str:
    from pypdf import PdfReader

    reader = PdfReader(io.BytesIO(data))
    if reader.is_encrypted:
        raise ValueError("encrypted PDF")
    lines: list[str] = []
    for n, page in enumerate(reader.pages, start=1):
        lines.append(f"[page {n}]")
        lines.append(page.extract_text() or "")
    return "\n".join(lines)


EXTRACTORS: Final[dict[str, Callable[[bytes], str]]] = {
    ".docx": _docx,
    ".xlsx": _xlsx,
    ".xlsm": _xlsx,  # values only; macros are never run
    ".pptx": _pptx,
    ".pdf": _pdf,
}


def extractable(name: str) -> bool:
    return PurePosixPath(name.replace("\\", "/")).suffix.lower() in EXTRACTORS


def extract_text(name: str, data: bytes) -> str | None:
    """Plain text of an Office/PDF file, or ``None`` if unsupported or unreadable."""
    extractor = EXTRACTORS.get(PurePosixPath(name.replace("\\", "/")).suffix.lower())
    if extractor is None or len(data) > MAX_DOCUMENT_BYTES:
        return None
    try:
        text = extractor(data)
    except Exception as exc:  # damaged, encrypted or unsupported variants are skipped
        log.info("extract_skipped", file=name, error=type(exc).__name__)
        return None
    return text[:MAX_EXTRACTED_CHARS]

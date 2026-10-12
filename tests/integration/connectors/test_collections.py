"""Knowledge collections (M8-04, FR-019/020, BR-012): folders, websites, Drive, OneDrive."""

import io
import re
from pathlib import Path
from typing import Any

import httpx
import pytest
from docx import Document as DocxDocument

from panos_core.config import PanosSettings
from panos_core.domain.connectors import CollectionKind, ConnectorKind
from panos_core.domain.enums import PrivacyMode
from panos_core.knowledge import KnowledgeIndex, KnowledgeService
from panos_core.knowledge.collections import (
    TIER_COMPANY,
    TIER_EXTERNAL,
    TIER_PROJECT,
    CollectionError,
    CollectionService,
)
from panos_core.knowledge.index import Document
from panos_core.knowledge.sources import html_to_text
from panos_core.projects import ProjectService
from panos_core.tools import ToolContext
from panos_core.tools.builtin.desktop import KnowledgeSearchInput, KnowledgeSearchTool
from tests.connector_fakes import query
from tests.integration.connectors.conftest import Env, connect_google

pytestmark = pytest.mark.req("FR-019", "FR-020", "BR-012")


def _docx(text: str) -> bytes:
    doc = DocxDocument()
    doc.add_paragraph(text)
    buffer = io.BytesIO()
    doc.save(buffer)
    return buffer.getvalue()


PAGES = {
    "/robots.txt": (200, "text/plain", b"User-agent: *\nDisallow: /private/\n"),
    "/docs/refunds": (
        200,
        "text/html; charset=utf-8",
        b"<html><head><title>Refunds</title><script>steal()</script></head><body>"
        b"<nav>Menu</nav><h1>Refund policy</h1><p>Refunds within 14 days.</p>"
        b"<form action='/submit'><input name='q'></form></body></html>",
    ),
    "/private/secret": (200, "text/html", b"<p>hidden</p>"),
    "/moved": (301, "text/html", b""),
    "/away": (302, "text/html", b""),
    "/big": (200, "text/plain", b"x" * 6_000_000),
}


class Web:
    def __init__(self) -> None:
        self.requests: list[httpx.Request] = []

    def handle(self, request: httpx.Request) -> httpx.Response:
        self.requests.append(request)
        if request.url.host == "down.example":
            raise httpx.ConnectError("down")
        status, ctype, body = PAGES.get(request.url.path, (404, "text/plain", b"missing"))
        headers = {"content-type": ctype}
        if request.url.path == "/moved":
            headers["Location"] = "/docs/refunds"
        if request.url.path == "/away":
            headers["Location"] = "https://elsewhere.example/x"
        return httpx.Response(status, headers=headers, content=body)


@pytest.fixture
def web() -> Web:
    return Web()


@pytest.fixture
def svc(env: Env, tmp_path: Path, web: Web) -> CollectionService:
    knowledge = KnowledgeService(KnowledgeIndex(tmp_path / "knowledge.db"), env.audit)
    return CollectionService(
        env.uow_factory,
        knowledge,
        env.service,
        env.audit,
        web_transport=httpx.MockTransport(web.handle),
    )


def test_html_to_text_drops_scripts_and_markup() -> None:
    title, text = html_to_text(PAGES["/docs/refunds"][2].decode())
    assert title == "Refunds"
    assert "steal" not in text and "<" not in text
    assert "Refund policy\nRefunds within 14 days." in text


async def test_folder_collection_is_incremental_and_ranked_after_the_project(
    svc: CollectionService, tmp_path: Path
) -> None:
    folder = tmp_path / "company"
    folder.mkdir()
    (folder / "handbook.docx").write_bytes(_docx("Holiday policy: 25 days of annual leave."))
    (folder / "notes.md").write_text("Annual leave requests go to HR.", encoding="utf-8")
    created = await svc.create("company", CollectionKind.FOLDER, {"path": str(folder)})
    assert created.privacy_mode == PrivacyMode.LOCAL_ONLY  # default (ADR-024)

    first = await svc.index("company")
    assert (first.stats.indexed, first.stats.files_seen) == (2, 2)
    again = await svc.index(created.id)
    assert (again.stats.indexed, again.stats.unchanged) == (0, 2)
    (folder / "notes.md").unlink()
    gone = await svc.index("company")
    assert gone.stats.removed == 1

    # The current project outranks company knowledge, whatever the scores (FR-020).
    index = svc._knowledge.index
    doc = Document("README.md", "Annual leave is not tracked here.", "h", 10)
    index.update("prj_shop", "STANDARD", [doc], {"README.md"})
    hits = await svc.search("annual leave", project_id="prj_shop")
    assert [(h.tier, h.source) for h in hits] == [
        (TIER_PROJECT, "project"),
        (TIER_COMPANY, "company"),
    ]
    assert hits[1].hit.uri == "handbook.docx" and hits[1].hit.privacy_mode == "LOCAL_ONLY"
    assert await svc.search("annual leave", collections=[]) == []

    await svc.update("company", privacy=PrivacyMode.STANDARD)
    [hit] = await svc.search("annual leave")
    assert hit.hit.privacy_mode == "STANDARD"  # relabelled in the index
    await svc.update("company", enabled=False)
    assert await svc.search("annual leave") == []
    with pytest.raises(CollectionError) as disabled:
        await svc.index("company")
    assert disabled.value.reason == "invalid_state"


async def test_invalid_collections_are_rejected(svc: CollectionService, tmp_path: Path) -> None:
    cases: list[tuple[str, CollectionKind, dict[str, Any]]] = [
        ("missing", CollectionKind.FOLDER, {"path": str(tmp_path / "nope")}),
        ("relative", CollectionKind.FOLDER, {"path": "docs"}),
        ("ftp", CollectionKind.WEBSITE, {"urls": ["ftp://files.example/x"]}),
        ("none", CollectionKind.WEBSITE, {"urls": []}),
        ("inject", CollectionKind.GOOGLE_DRIVE, {"folder_id": "x' or name contains '"}),
        ("bad name!", CollectionKind.WEBSITE, {"urls": ["https://a.example"]}),
    ]
    for name, kind, location in cases:
        with pytest.raises(CollectionError):
            await svc.create(name, kind, location)
    with pytest.raises(CollectionError) as no_account:
        await svc.create("drive", CollectionKind.GOOGLE_DRIVE, {"folder_id": "abc"})
    assert "connect" in str(no_account.value)
    await svc.create("site", CollectionKind.WEBSITE, {"urls": ["https://docs.example/a"]})
    with pytest.raises(CollectionError) as dup:
        await svc.create("site", CollectionKind.WEBSITE, {"urls": ["https://docs.example/b"]})
    assert dup.value.reason == "duplicate"


async def test_websites_honour_robots_redirects_and_size_limits(
    svc: CollectionService, web: Web
) -> None:
    urls = [
        "https://docs.example/docs/refunds",
        "https://docs.example/private/secret",
        "https://docs.example/moved",
        "https://docs.example/away",
        "https://docs.example/big",
        "https://down.example/page",
    ]
    await svc.create("vendor-docs", CollectionKind.WEBSITE, {"urls": urls})
    result = await svc.index("vendor-docs")
    assert result.stats.indexed == 2  # the page itself and the same-site redirect to it
    reasons = "\n".join(result.skipped)
    assert "private/secret: disallowed by robots.txt" in reasons
    assert "away: redirects to another site" in reasons
    assert "big: too large" in reasons
    assert "down.example/page: failed" in reasons
    assert all(r.method == "GET" for r in web.requests)  # never submits the form
    assert not any(r.url.path == "/submit" for r in web.requests)
    assert not any(r.url.host == "elsewhere.example" for r in web.requests)
    [hit] = [h for h in await svc.search("refund") if h.hit.uri.endswith("/docs/refunds")]
    assert hit.tier == TIER_EXTERNAL and hit.source == "vendor-docs"
    assert "steal" not in hit.hit.snippet


async def test_google_drive_folder_is_read_only_and_incremental(
    svc: CollectionService, env: Env
) -> None:
    files = {
        "d1": {
            "id": "d1",
            "name": "Pricing",
            "mimeType": "application/vnd.google-apps.document",
            "modifiedTime": "2026-10-01T10:00:00Z",
            "webViewLink": "https://docs.google.com/document/d/d1/edit",
        },
        "sub": {"id": "sub", "name": "Contracts", "mimeType": "application/vnd.google-apps.folder"},
        "w1": {
            "id": "w1",
            "name": "nda.docx",
            "mimeType": "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
            "modifiedTime": "2026-10-02T10:00:00Z",
            "md5Checksum": "abc",
            "size": "2000",
            "webViewLink": "https://drive.google.com/file/d/w1/view",
        },
    }
    parents = {"root1": ["d1", "sub"], "sub": ["w1"]}

    def listing(request: httpx.Request, m: re.Match[str]) -> httpx.Response:
        folder = re.match(r"'([^']+)' in parents", query(request)["q"])
        ids = parents.get(folder.group(1) if folder else "", [])
        return httpx.Response(200, json={"files": [files[i] for i in ids]})

    def export(request: httpx.Request, m: re.Match[str]) -> httpx.Response:
        assert query(request)["mimeType"] == "text/plain"
        return httpx.Response(200, text="Enterprise plan costs 99 EUR per seat.")

    def media(request: httpx.Request, m: re.Match[str]) -> httpx.Response:
        assert query(request)["alt"] == "media"
        return httpx.Response(200, content=_docx("Confidentiality lasts five years."))

    env.cloud.routes += [
        ("GET", r"/drive/v3/files", listing),
        ("GET", r"/drive/v3/files/([^/]+)/export", export),
        ("GET", r"/drive/v3/files/([^/]+)", media),
    ]
    await connect_google(env, ["drive"])
    await svc.create("sales-drive", CollectionKind.GOOGLE_DRIVE, {"folder_id": "root1"})
    first = await svc.index("sales-drive")
    assert first.stats.indexed == 2, first.skipped
    downloads = len(env.cloud.calls("GET", "/export")) + len(env.cloud.calls("GET", "/files/w1"))
    again = await svc.index("sales-drive")
    assert again.stats.unchanged == 2
    assert len(env.cloud.calls("GET", "/export")) + len(env.cloud.calls("GET", "/files/w1")) == (
        downloads
    )  # unchanged remote versions are not downloaded again
    assert {r.method for r in env.cloud.requests if "/drive/" in r.url.path} == {"GET"}
    [hit] = await svc.search("confidentiality")
    assert hit.hit.uri == "https://drive.google.com/file/d/w1/view"  # provenance link


async def test_onedrive_follows_download_redirect_without_the_token(
    svc: CollectionService, env: Env
) -> None:
    children = {
        "value": [
            {
                "id": "f1",
                "name": "Onboarding.md",
                "file": {},
                "size": 40,
                "eTag": "e1",
                "webUrl": "https://contoso.sharepoint.com/Shared/Onboarding.md",
            },
            {"id": "fold", "name": "Archive", "folder": {"childCount": 0}},
        ]
    }

    def listing(request: httpx.Request, m: re.Match[str]) -> httpx.Response:
        return httpx.Response(200, json=children if m.group(1) == "root" else {"value": []})

    def content(request: httpx.Request, m: re.Match[str]) -> httpx.Response:
        return httpx.Response(302, headers={"Location": "https://dl.contoso.example/dl/f1-blob"})

    def download(request: httpx.Request) -> httpx.Response:
        assert "Authorization" not in request.headers  # pre-authenticated URL only
        return httpx.Response(200, text="New staff get a laptop on day one.")

    env.cloud.routes += [
        ("GET", r"/v1.0/drives/d-1/items/([^/]+)/children", listing),
        ("GET", r"/v1.0/drives/d-1/items/([^/]+)/content", content),
        ("GET", r"/dl/f1-blob", lambda request, m: download(request)),
    ]
    env.service.set_client(ConnectorKind.MICROSOFT, "ms-client")
    flow = await env.service.start_microsoft()
    assert flow.task is not None
    await flow.task
    await svc.create("hr", CollectionKind.ONEDRIVE, {"drive_id": "d-1"})
    result = await svc.index("hr")
    assert result.stats.indexed == 1, result.skipped
    assert env.cloud.calls("GET", "dl.contoso.example")
    [hit] = await svc.search("laptop")
    assert hit.hit.uri == "https://contoso.sharepoint.com/Shared/Onboarding.md"
    await svc.delete("hr")
    assert await svc.search("laptop") == []


async def test_knowledge_search_tool_ranks_project_before_collections(
    svc: CollectionService, env: Env, tmp_path: Path, settings: PanosSettings
) -> None:
    folder = tmp_path / "policies"
    folder.mkdir()
    (folder / "leave.md").write_text("Annual leave is 25 days.", encoding="utf-8")
    await svc.create("policies", CollectionKind.FOLDER, {"path": str(folder)})
    await svc.index("policies")
    (tmp_path / "shop").mkdir()
    project = await ProjectService(env.uow_factory, env.audit, settings, env.node.id).register(
        tmp_path / "shop"
    )
    doc = Document("README.md", "Annual leave sync job.", "h", 10)
    svc._knowledge.index.update(project.id, "LOCAL_ONLY", [doc], {"README.md"})
    tool = KnowledgeSearchTool(svc._knowledge, svc)
    ctx = ToolContext(actor="user", node_id=env.node.id, project=project)
    inp = KnowledgeSearchInput(query="annual leave")
    out = await tool._execute(inp, ctx, await tool.plan(inp, ctx))
    assert [(h.source, h.tier, h.uri) for h in out.hits] == [
        ("project", 1, "README.md"),
        ("policies", 2, "leave.md"),
    ]
    only = KnowledgeSearchInput(query="annual leave", collections=[])
    assert [
        h.source for h in (await tool._execute(only, ctx, await tool.plan(only, ctx))).hits
    ] == ["project"]
    everything = KnowledgeSearchInput(query="annual leave", all_projects=True)
    hits = (await tool._execute(everything, ctx, await tool.plan(everything, ctx))).hits
    assert {h.source for h in hits} == {"index"} and len(hits) == 2

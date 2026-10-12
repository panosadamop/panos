"""Office/PDF documents and Git history in the knowledge index (M8-04, FR-019)."""

import io
import shutil
import subprocess
from pathlib import Path

import pytest
from docx import Document as DocxDocument
from openpyxl import Workbook
from pptx import Presentation

from panos_core.documents import FontNotFoundError, find_fonts, render_pdf
from panos_core.knowledge.extract import extract_text
from panos_core.knowledge.git_history import URI as HISTORY
from tests.integration.projects.conftest import M3
from tests.repos import copy_repo

pytestmark = pytest.mark.req("FR-019", "BR-012", "NFR-PRV-003")


def _docx() -> bytes:
    doc = DocxDocument()
    doc.add_paragraph("Η πολιτική επιστροφών ισχύει για 30 ημέρες.")
    table = doc.add_table(rows=2, cols=2)
    table.rows[0].cells[0].text, table.rows[0].cells[1].text = "SLA", "Value"
    table.rows[1].cells[0].text, table.rows[1].cells[1].text = "refund window", "30 days"
    buffer = io.BytesIO()
    doc.save(buffer)
    return buffer.getvalue()


def _xlsx() -> bytes:
    wb = Workbook()
    ws = wb.active
    assert ws is not None
    ws.title = "Rates"
    ws.append(["Region", "Shipping"])
    ws.append(["Crete", 7.5])
    buffer = io.BytesIO()
    wb.save(buffer)
    return buffer.getvalue()


def _pptx() -> bytes:
    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[1])
    slide.shapes.title.text = "Roadmap"
    slide.placeholders[1].text_frame.text = "Loyalty programme in Q3"
    slide.notes_slide.notes_text_frame.text = "Ask finance about the budget"
    buffer = io.BytesIO()
    prs.save(buffer)
    return buffer.getvalue()


def test_extractors_keep_location_markers_and_skip_damaged_files() -> None:
    docx = extract_text("policy.docx", _docx())
    assert docx is not None and "επιστροφών" in docx and "refund window | 30 days" in docx
    xlsx = extract_text("rates.xlsx", _xlsx())
    assert xlsx == "[sheet Rates]\nRegion\tShipping\nCrete\t7.5"
    pptx = extract_text("plan.pptx", _pptx())
    assert pptx is not None
    assert "[slide 1]" in pptx and "Loyalty programme in Q3" in pptx
    assert "[notes] Ask finance about the budget" in pptx
    assert extract_text("broken.docx", b"not a zip") is None
    assert extract_text("notes.txt", b"plain") is None  # text files don't need an extractor
    try:
        pdf = render_pdf("# Manual\n\nWarranty is two years.\n", title="Manual", fonts=find_fonts())
    except FontNotFoundError:
        return
    text = extract_text("manual.pdf", pdf)
    assert text is not None and text.startswith("[page 1]") and "Warranty is two years." in text


async def test_office_documents_and_git_history_are_searchable(m3: M3, tmp_path: Path) -> None:
    repo = copy_repo("python_fastapi", tmp_path, git=True)
    if not shutil.which("git"):
        pytest.skip("git not installed")
    docs = repo / "docs"
    docs.mkdir(exist_ok=True)
    (docs / "policy.docx").write_bytes(_docx())
    (docs / "rates.xlsx").write_bytes(_xlsx())
    (docs / "roadmap.pptx").write_bytes(_pptx())
    (docs / "~$policy.docx").write_bytes(b"\x00lock")  # Office lock file: skipped
    subprocess.run(
        [
            "git",
            "-C",
            str(repo),
            "-c",
            "user.email=t@example.invalid",
            "-c",
            "user.name=Tester",
            "commit",
            "-q",
            "--allow-empty",
            "-m",
            "Fix payment retry backoff",
            "-m",
            "Closes #42",
        ],
        check=True,
    )
    project = await m3.projects.register(repo)
    first = await m3.knowledge.index_project(project)
    assert first.indexed == first.files_seen

    for query, uri in [
        ("επιστροφών", "docs/policy.docx"),
        ("Crete", "docs/rates.xlsx"),
        ("loyalty", "docs/roadmap.pptx"),
        ("retry backoff", HISTORY),
    ]:
        hits = await m3.knowledge.search(query, project_id=project.id)
        assert hits and hits[0].uri == uri, (query, [h.uri for h in hits])

    again = await m3.knowledge.index_project(project)
    assert again.indexed == 0  # history re-indexed only when HEAD moves
    subprocess.run(
        [
            "git",
            "-C",
            str(repo),
            "-c",
            "user.email=t@example.invalid",
            "-c",
            "user.name=Tester",
            "commit",
            "-q",
            "--allow-empty",
            "-m",
            "Add invoice export",
        ],
        check=True,
    )
    moved = await m3.knowledge.index_project(project)
    assert moved.indexed == 1
    assert (await m3.knowledge.search("invoice export", project_id=project.id))[0].uri == HISTORY

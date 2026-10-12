"""Office/PDF exports (M8-05): open every generated file and check its structure."""

import io
from pathlib import Path

import pytest
from docx import Document
from openpyxl import load_workbook
from pptx import Presentation
from pypdf import PdfReader

from panos_core.documents import FontNotFoundError, FontSet, find_fonts, render_docx, render_pdf
from panos_core.documents.markdown import Code, Heading, ListBlock, Quote, Table, parse, runs
from panos_core.requirements.exports import ExportError, export
from tests.unit.requirements.test_artifacts_golden import shop

pytestmark = pytest.mark.req("FR-029", "BR-028")

GREEK = "# Έργο\n\nΗ πληρωμή με **κάρτα** γίνεται _μία_ φορά — `POST /orders`.\n"


def _fonts_or_skip() -> FontSet:
    try:
        return find_fonts()
    except FontNotFoundError:
        pytest.skip("no Unicode TrueType font on this machine")


def test_markdown_subset_is_parsed_into_blocks_and_runs() -> None:
    blocks = parse(
        "# T\n\n> note\n> more\n\n- a\n  - **b** x\n1. one\n\n"
        "| A | B |\n|---|---|\n| 1 | 2 \\| 3 |\n\n"
        "```mermaid\nerDiagram\n```\n"
    )
    assert isinstance(blocks[0], Heading) and blocks[0].text == "T"
    assert isinstance(blocks[1], Quote) and blocks[1].text == "note more"
    lists = blocks[2]
    assert isinstance(lists, ListBlock)
    assert [(i.depth, i.ordered) for i in lists.items] == [(0, False), (1, False), (0, True)]
    table = blocks[3]
    assert isinstance(table, Table) and table.rows == (("1", "2 | 3"),)
    assert isinstance(blocks[4], Code) and blocks[4].language == "mermaid"
    styled = [(r.text, r.bold, r.italic, r.code) for r in runs("**a *b*** `c` snake_case_name _d_")]
    assert styled == [
        ("a ", True, False, False),
        ("b", True, True, False),
        (" ", False, False, False),
        ("c", False, False, True),
        (" snake_case_name ", False, False, False),
        ("d", False, True, False),
    ]


def test_specification_docx_has_every_document_with_headings_tables_and_lists() -> None:
    files = export(shop(), ["docx"])
    doc = Document(io.BytesIO(files["SPECIFICATION.docx"]))
    titles = [p.text for p in doc.paragraphs if p.style and p.style.name == "Title"]
    assert titles[0] == "Web Shop — Business Requirements (BRD)"
    assert len(titles) == 11  # one per Markdown document, each on a new page
    assert any(p.style and p.style.name.startswith("List Bullet") for p in doc.paragraphs)
    estimates = [t for t in doc.tables if t.rows[0].cells[0].text == "Story"]
    assert estimates and estimates[0].rows[1].cells[0].text.startswith("US-001")
    text = "\n".join(p.text for p in doc.paragraphs)
    assert "**" not in text and "AC-US001-02" in text and "STALE" in text
    assert doc.core_properties.title == "Web Shop — Specification"


def test_specification_pdf_contains_greek_and_all_documents() -> None:
    fonts = _fonts_or_skip()
    pdf = PdfReader(io.BytesIO(export(shop(), ["pdf"], fonts=fonts)["SPECIFICATION.pdf"]))
    text = "\n".join(page.extract_text() for page in pdf.pages)
    assert len(pdf.pages) >= 11
    assert "Web Shop — Risks" in text and "US-001" in text
    greek = PdfReader(io.BytesIO(render_pdf(GREEK, title="Έργο", fonts=fonts)))
    assert "κάρτα" in greek.pages[0].extract_text()
    assert greek.metadata is not None and greek.metadata.title == "Έργο"


def test_workbook_sheets_are_filterable_and_complete() -> None:
    wb = load_workbook(io.BytesIO(export(shop(), ["xlsx"])["REQUIREMENTS.xlsx"]))
    assert wb.sheetnames == ["Requirements", "Backlog", "Cost", "Sprints", "Trace links"]
    reg = wb["Requirements"]
    assert [c.value for c in reg[1]][:5] == ["Key", "Type", "Title", "Version", "Status"]
    assert reg.max_row == 1 + len(shop().items)
    assert reg.freeze_panes == "A2" and reg.auto_filter.ref
    stale = [row for row in reg.iter_rows(min_row=2, values_only=True) if row[0] == "AC-US001-02"]
    assert stale[0][5] == "yes"
    backlog = list(wb["Backlog"].iter_rows(min_row=2, values_only=True))
    assert backlog[1][:4] == ("US-002", "As an admin I can add a product", 3, "default")
    assert list(wb["Trace links"].iter_rows(min_row=2, values_only=True)) == [
        ("API-001", "realizes", "FR-001")
    ]


def test_summary_deck_covers_scope_estimate_risks_and_open_decisions() -> None:
    prs = Presentation(io.BytesIO(export(shop(), ["pptx"])["SUMMARY.pptx"]))
    titles = [s.shapes.title.text for s in prs.slides]
    assert titles == [
        "Web Shop",
        "Business goals",
        "Scope",
        "Technology stack",
        "Estimate",
        "Risks",
        "Open decisions",
    ]
    open_slide = "\n".join(sh.text_frame.text for sh in prs.slides[6].shapes if sh.has_text_frame)
    assert "AC-US001-02" in open_slide and "STALE" in open_slide and "FR-002" in open_slide


def test_unknown_or_empty_formats_are_rejected() -> None:
    with pytest.raises(ExportError):
        export(shop(), ["odt"])
    with pytest.raises(ExportError):
        export(shop(), [])


def test_docx_renders_greek_inline_styles() -> None:
    doc = Document(io.BytesIO(render_docx(GREEK, title="Έργο")))
    body = doc.paragraphs[1]
    assert body.text == "Η πληρωμή με κάρτα γίνεται μία φορά — POST /orders."
    assert [r.text for r in body.runs if r.bold] == ["κάρτα"]
    assert [r.text for r in body.runs if r.italic] == ["μία"]


def test_pdf_font_lookup_reports_a_clear_error(tmp_path: Path) -> None:
    with pytest.raises(FontNotFoundError, match="pdf_font_dir"):
        find_fonts(tmp_path)
